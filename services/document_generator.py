"""Document Generator for Belarus OHS forms:
1. Act Form H-1 (Акт о несчастном случае на производстве формы Н-1).
2. Work Permit (Наряд-допуск на выполнение работ повышенной опасности).

Provides structured data models, regulatory auto-enrichment against BY NPA,
and DOCX export with official Belarusian styling.
"""

from datetime import date
from io import BytesIO
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pydantic import BaseModel, Field


# ============================================================
# PYDANTIC SCHEMAS
# ============================================================

class ActN1Request(BaseModel):
    organization_name: str = Field(..., description="Наименование организации / работодателя")
    subdivision: str = Field("Производственный участок", description="Структурное подразделение / цех")
    accident_date: str = Field(default_factory=lambda: date.today().strftime("%d.%m.%Y"), description="Дата происшествия (ДД.ММ.ГГГГ)")
    accident_time: str = Field("10:00", description="Время происшествия (ЧЧ:ММ)")
    accident_location: str = Field(..., description="Место происшествия (оборудование, площадка, рабочее место)")
    victim_fio: str = Field(..., description="ФИО потерпевшего")
    victim_profession: str = Field(..., description="Профессия / должность потерпевшего")
    victim_age: Optional[int] = Field(None, description="Возраст потерпевшего")
    experience_total: str = Field("5 лет", description="Общий стаж работы")
    experience_profession: str = Field("2 года", description="Стаж работы по профессии / должности")
    briefing_date: Optional[str] = Field(None, description="Дата проведения инструктажа по охране труда")
    knowledge_check_date: Optional[str] = Field(None, description="Дата проверки знаний по вопросам охраны труда")
    incident_description: str = Field(..., description="Обстоятельства несчастного случая (фабула)")
    causes: Optional[List[str]] = Field(None, description="Причины несчастного случая (основная и сопутствующие)")
    violated_npa: Optional[List[str]] = Field(None, description="Нарушенные требования НПА и локальных актов")
    responsible_persons: Optional[List[str]] = Field(None, description="Лица, допустившие нарушения актов законодательства")
    corrective_measures: Optional[List[Dict[str, str]]] = Field(None, description="Мероприятия по устранению причин (мероприятие, срок, ответственный)")


class WorkPermitRequest(BaseModel):
    organization_name: str = Field(..., description="Наименование организации")
    subdivision: str = Field("Цех № 1", description="Подразделение / участок")
    work_type: str = Field(..., description="Характер работ (на высоте, огневые, в замкнутом пространстве и др.)")
    work_location: str = Field(..., description="Место проведения работ")
    start_date: str = Field(default_factory=lambda: date.today().strftime("%d.%m.%Y"), description="Дата начала")
    start_time: str = Field("08:00", description="Время начала")
    end_date: str = Field(default_factory=lambda: date.today().strftime("%d.%m.%Y"), description="Дата окончания")
    end_time: str = Field("17:00", description="Время окончания")
    permit_issuer: str = Field(..., description="Лицо, выдавшее наряд-допуск (должность, ФИО)")
    work_leader: str = Field(..., description="Ответственный руководитель работ (должность, ФИО)")
    work_supervisor: str = Field(..., description="Ответственный производитель работ (должность, ФИО)")
    brigade_members: List[Dict[str, str]] = Field(
        default_factory=list,
        description="Состав бригады: [{'fio': 'Иванов И.И.', 'qualification': '3 разряд / 1 группа на высоте'}]",
    )
    safety_measures: Optional[List[str]] = Field(None, description="Меры безопасности до начала и во время работ")
    required_ppe: Optional[List[str]] = Field(None, description="Перечень обязательных СИЗ")


# ============================================================
# REGULATORY AUTO-ENRICHMENT LOGIC
# ============================================================

def auto_suggest_act_n1_rules(incident_description: str) -> Dict[str, Any]:
    """Suggest likely violated Belarusian rules and causes based on incident text."""
    desc = incident_description.lower()
    violated: List[str] = []
    causes: List[str] = []
    measures: List[Dict[str, str]] = []

    # Height work / fall incidents
    if any(k in desc for k in ("высот", "лестниц", "паден", "стремянк", "леса", "подмост")):
        violated.extend([
            "пп. 120, 124 Правил по охране труда при выполнении работ на высоте (постановление Минтруда от 06.02.2025 № 11) — неприменение или ненадлежащее применение страховочной системы (предельное усилие не более 6 кН, длина стропа не более 2 м)",
            "п. 54 Правил по охране труда при выполнении работ на высоте № 11 — использование лестниц без предварительного визуального осмотра и проверки состояния",
            "ст. 17 Закона Республики Беларусь «Об охране труда» — необеспечение безопасных условий труда при организации работ на высоте",
        ])
        causes.extend([
            "Неприменение работающим выданных средств индивидуальной защиты от падения с высоты (страховочной привязи, стропа с амортизатором)",
            "Недостаточный контроль со стороны должностных лиц за соблюдением работающими требований безопасности на высоте",
        ])
        measures.append({
            "measure": "Провести внеплановый инструктаж по Правилам охраны труда на высоте № 11 со всеми работниками подразделения",
            "deadline": "В течение 3 рабочих дней",
            "responsible": "Начальник подразделения",
        })

    # Machine / lathe / mechanical equipment
    elif any(k in desc for k in ("станок", "токарн", "заготовк", "патрон", "вращающ", "шпиндель")):
        violated.extend([
            "Правила по охране труда при холодной обработке металлов — работа на оборудовании без защитного экрана зоны обработки заготовок",
            "Инструкция по охране труда для станочника — не отключение станка при установке и измерении детали",
            "ст. 17 Закона Республики Беларусь «Об охране труда» — допуск к эксплуатации оборудования с неисправными защитными ограждениями",
        ])
        causes.extend([
            "Эксплуатация токарного станка с открытой рабочей зоной без защитного кожуха/экрана",
            "Нарушение работающим технологической дисциплины и требований инструкции по охране труда",
        ])
        measures.append({
            "measure": "Оснастить станки защитными экранами, сблокированными с пуском, проверить исправность блокировок",
            "deadline": "В течение 5 рабочих дней",
            "responsible": "Главный механик",
        })

    # Electrical incidents
    elif any(k in desc for k in ("электро", "ток", "напряжен", "кабел", "заземлен", "щит")):
        violated.extend([
            "ТКП 181-2009 (Правила технической эксплуатации электроустановок потребителей) — невыполнение технических мероприятий, обеспечивающих безопасность работ со снятием напряжения",
            "ТКП 290-2010 (Правила применения и испытания средств защиты в электроустановках) — неприменение диэлектрических средств защиты",
        ])
        causes.extend([
            "Прикосновение к токоведущим частям, находящимся под напряжением, вследствие отсутствия защитного заземления/ограждения",
        ])
        measures.append({
            "measure": "Провести внеочередную проверку знаний электротехнического персонала по электробезопасности",
            "deadline": "В течение 10 рабочих дней",
            "responsible": "Лицо, ответственное за электрохозяйство",
        })

    # Fallback default OHS rules
    else:
        violated.extend([
            "ст. 17 Закона Республики Беларусь «Об охране труда» — ненадлежащий контроль за соблюдением требований безопасности на рабочем месте",
            "ст. 19 Закона Республики Беларусь «Об охране труда» — нарушение работающим правил и инструкций по охране труда",
        ])
        causes.append("Нарушение работающим требований локальных правовых актов по охране труда")
        measures.append({
            "measure": "Провести внеплановый инструктаж по охране труда с персоналом участка",
            "deadline": "В течение 3 дней",
            "responsible": "Мастер участка",
        })

    return {
        "violated_npa": violated,
        "causes": causes,
        "corrective_measures": measures,
    }


def auto_suggest_work_permit_measures(work_type: str) -> Dict[str, Any]:
    """Suggest required safety measures and PPE for work permit by work category."""
    wt = work_type.lower()
    measures: List[str] = []
    ppe: List[str] = [
        "Каска защитная с подбородочным ремнем",
        "Специальная одежда и специальная обувь",
        "Защитные перчатки",
    ]

    if any(k in wt for k in ("высот", "верхолаз", "кровл", "монтаж")):
        measures.extend([
            "Определить места и схемы закрепления страховочных и анкерных систем согласно ППР",
            "Проверить наличие и исправность страховочных привязей, стропов с амортизаторами и карабинов",
            "Оградить зону возможного падения предметов снизу сигнальной лентой и выставить предупреждающие знаки «Опасная зона»",
            "Проверить устойчивость и целостность строительных лесов/подмостей/стремянок с составлением акта приемки",
            "Запретить производство работ при скорости ветра 15 м/с и более, грозе, гололеде",
        ])
        ppe.extend([
            "Страховочная привязь (ГОСТ EN 361)",
            "Строп с амортизатором (длина не более 2 м, предельное усилие не более 6 кН)",
            "Карабины самозапирающиеся с предохранительной муфтой",
        ])

    elif any(k in wt for k in ("огнев", "сварк", "резк", "пайк")):
        measures.extend([
            "Очистить место проведения огневых работ от горючих материалов в радиусе не менее 5 м",
            "Оснастить место проведения работ не менее чем двумя порошковыми или углекислотными огнетушителями",
            "Полить водой прилегающие сгораемые конструкции и защитить их негорючими экранами",
            "Обеспечить контроль за местом проведения огневых работ в течение не менее 3 часов после их окончания",
            "Проверить отсутствие взрывоопасных концентраций паров и газов газоанализатором",
        ])
        ppe.extend([
            "Костюм сварщика термостойкий",
            "Щиток защитный лицевой с защитным светофильтром",
            "Краги спилковые",
        ])

    elif any(k in wt for k in ("замкнут", "колод", "емкост", "резервуар", "отсек")):
        measures.extend([
            "Произвести предварительную вентиляцию (проветривание) замкнутого пространства",
            "Выполнить замер состава воздушной среды переносным газоанализатором на кислород, угарный газ и метан",
            "Обеспечить непрерывное дежурство двух страхующих работников у лаза/люка емкости",
            "Оснастить работающего внутри спасательной привязью со спасательным канатом, выведенным наружу",
            "Использовать для освещения переносные светильники напряжением не выше 12 В во взрывозащищенном исполнении",
        ])
        ppe.extend([
            "Спасательная привязь с наплечными и набедренными лямками",
            "Шланговый противогаз или изолирующий дыхательный аппарат",
            "Газоанализатор непрерывного действия индивидуальный",
        ])

    else:
        measures.extend([
            "Провести целевой инструктаж всем членам бригады с записью в наряде-допуске",
            "Оградить опасную зону сигнальным ограждением",
            "Проверить исправность рабочего инструмента и защитных блокировок",
        ])

    return {
        "safety_measures": measures,
        "required_ppe": list(dict.fromkeys(ppe)),
    }


# ============================================================
# DOCX FORMATTING HELPERS
# ============================================================

def _font(run, size=10, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _p(doc, text="", size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        _font(r, size=size, bold=bold, italic=italic)
    return p


def _cell(cell, text="", size=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text or ""))
    _font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_cell_width(cell, cm: float):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(cm * 567)))
    tcW.set(qn("w:type"), "dxa")


def _table_borders(table, color="000000", size="4"):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = borders.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge)
            borders.append(e)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), size)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)


# ============================================================
# ACT FORM H-1 GENERATOR
# ============================================================

def generate_act_n1_docx(req: ActN1Request) -> BytesIO:
    """Generate official Belarus Act Form H-1 as DOCX."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(1.5)

    # УТВЕРЖДАЮ
    p_app = _p(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
    r = p_app.add_run("УТВЕРЖДАЮ\nРуководитель организации\n")
    _font(r, size=11, bold=True)
    r2 = p_app.add_run(f"________________ {req.organization_name}\n«___» ____________ 202_ г.\nМ.П.\n")
    _font(r2, size=10)

    # Title
    _p(doc, "АКТ № ______", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(doc, "о несчастном случае на производстве (форма Н-1)", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # Section 1
    _p(doc, f"1. Организация (наниматель): {req.organization_name}", size=10, bold=True)
    _p(doc, f"   Структурное подразделение: {req.subdivision}", size=10)

    # Section 2
    _p(doc, "2. Сведения о потерпевшем:", size=10, bold=True)
    _p(doc, f"   Фамилия, собственное имя, отчество: {req.victim_fio}", size=10)
    _p(doc, f"   Профессия (должность): {req.victim_profession}", size=10)
    if req.victim_age:
        _p(doc, f"   Возраст: {req.victim_age} лет", size=10)
    _p(doc, f"   Стаж работы общий: {req.experience_total}, по профессии: {req.experience_profession}", size=10)
    if req.briefing_date:
        _p(doc, f"   Дата проведения инструктажа по охране труда: {req.briefing_date}", size=10)
    if req.knowledge_check_date:
        _p(doc, f"   Дата проверки знаний по вопросам охраны труда: {req.knowledge_check_date}", size=10)

    # Section 3
    _p(doc, f"3. Дата и время несчастного случая: {req.accident_date} в {req.accident_time}", size=10, bold=True)
    _p(doc, f"   Место несчастного случая: {req.accident_location}", size=10)

    # Section 4: Circumstances
    _p(doc, "4. Обстоятельства несчастного случая:", size=10, bold=True)
    _p(doc, req.incident_description, size=10, space_after=6)

    # Auto-enrich if not provided
    auto = auto_suggest_act_n1_rules(req.incident_description)
    causes = req.causes or auto["causes"]
    violated = req.violated_npa or auto["violated_npa"]
    measures = req.corrective_measures or auto["corrective_measures"]

    # Section 5: Causes
    _p(doc, "5. Причины несчастного случая:", size=10, bold=True)
    for idx, cause in enumerate(causes, 1):
        _p(doc, f"   5.{idx}. {cause}", size=10)

    # Section 6: Violated rules
    _p(doc, "6. Нарушенные акты законодательства, локальные правовые акты:", size=10, bold=True)
    for idx, rule in enumerate(violated, 1):
        _p(doc, f"   6.{idx}. {rule}", size=9, italic=True)

    # Section 7: Responsible persons
    _p(doc, "7. Лица, допустившие нарушения требований по охране труда:", size=10, bold=True)
    if req.responsible_persons:
        for p in req.responsible_persons:
            _p(doc, f"   - {p}", size=10)
    else:
        _p(doc, "   Лица, допустившие нарушения, устанавливаются комиссией по расследованию.", size=10)

    # Section 8: Corrective measures table
    _p(doc, "8. Мероприятия по устранению причин несчастного случая:", size=10, bold=True, space_after=4)
    table = doc.add_table(rows=1, cols=4)
    _table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["№ п/п", "Наименование мероприятия", "Срок исполнения", "Ответственный"]
    widths = [1.2, 8.5, 3.2, 4.0]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        _cell(hdr_cells[i], h, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_width(hdr_cells[i], widths[i])

    for i, m in enumerate(measures, 1):
        row = table.add_row().cells
        _cell(row[0], str(i), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_width(row[0], widths[0])
        _cell(row[1], m.get("measure", ""), size=9)
        _set_cell_width(row[1], widths[1])
        _cell(row[2], m.get("deadline", ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_width(row[2], widths[2])
        _cell(row[3], m.get("responsible", ""), size=9)
        _set_cell_width(row[3], widths[3])

    _p(doc, "", space_after=12)

    # Signatures
    _p(doc, "Лица, проводившие расследование несчастного случая:", size=10, bold=True)
    _p(doc, "Руководитель службы охраны труда (специалист по ОТ): _________ / _________________ /", size=10)
    _p(doc, "Представитель профсоюза (уполномоченное лицо): _________ / _________________ /", size=10)
    _p(doc, "Руководитель подразделения: _________ / _________________ /", size=10)

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out


# ============================================================
# WORK PERMIT (НАРЯД-ДОПУСК) GENERATOR
# ============================================================

def generate_work_permit_docx(req: WorkPermitRequest) -> BytesIO:
    """Generate official Belarus Work Permit (Наряд-допуск) as DOCX."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(1.5)

    _p(doc, req.organization_name, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    _p(doc, req.subdivision, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    _p(doc, "НАРЯД-ДОПУСК № ______", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(doc, "на выполнение работ с повышенной опасностью", size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    _p(doc, f"1. Выдан (кому): Ответственному руководителю работ: {req.work_leader}", size=10)
    _p(doc, f"   Ответственному производителю работ: {req.work_supervisor}", size=10)
    _p(doc, f"2. На выполнение работ: {req.work_type}", size=10, bold=True)
    _p(doc, f"3. Место проведения работ: {req.work_location}", size=10)
    _p(doc, f"4. Сроки проведения работ: с {req.start_time} {req.start_date} по {req.end_time} {req.end_date}", size=10)

    # Auto-enrich measures and PPE if not provided
    auto = auto_suggest_work_permit_measures(req.work_type)
    measures = req.safety_measures or auto["safety_measures"]
    ppe = req.required_ppe or auto["required_ppe"]

    # Safety measures
    _p(doc, "5. Мероприятия по подготовке рабочего места и обеспечению безопасности:", size=10, bold=True, space_after=3)
    for idx, m in enumerate(measures, 1):
        _p(doc, f"   5.{idx}. {m}", size=9)

    # Required PPE
    _p(doc, "6. Необходимые средства индивидуальной защиты (СИЗ):", size=10, bold=True, space_after=3)
    for p in ppe:
        _p(doc, f"   • {p}", size=9)

    # Brigade members table
    _p(doc, "7. Состав бригады и отметка о проведении целевого инструктажа:", size=10, bold=True, space_after=4)
    table = doc.add_table(rows=1, cols=5)
    _table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["№", "Ф.И.О. члена бригады", "Профессия / разряд / группа", "Подпись прошедшего инструктаж", "Дата"]
    widths = [1.0, 5.5, 4.5, 3.5, 2.5]
    for i, h in enumerate(headers):
        _cell(table.rows[0].cells[i], h, size=8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_width(table.rows[0].cells[i], widths[i])

    members = req.brigade_members or [
        {"fio": req.work_supervisor, "qualification": "Производитель работ"},
        {"fio": "Работник 1", "qualification": "Монтажник 4 р. / 1 группа"},
        {"fio": "Работник 2", "qualification": "Монтажник 3 р. / 1 группа"},
    ]

    for i, mem in enumerate(members, 1):
        row = table.add_row().cells
        _cell(row[0], str(i), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_width(row[0], widths[0])
        _cell(row[1], mem.get("fio", ""), size=8)
        _set_cell_width(row[1], widths[1])
        _cell(row[2], mem.get("qualification", ""), size=8)
        _set_cell_width(row[2], widths[2])
        _cell(row[3], "", size=8)
        _set_cell_width(row[3], widths[3])
        _cell(row[4], req.start_date, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_width(row[4], widths[4])

    _p(doc, "", space_after=10)

    # Signatures
    _p(doc, f"8. Наряд-допуск выдал: {req.permit_issuer}  Подпись: ___________  Дата: {req.start_date}", size=9)
    _p(doc, "9. Мероприятия выполнены, рабочее место подготовлено.", size=9, bold=True)
    _p(doc, f"   Ответственный руководитель работ: _________ / {req.work_leader} /  Подпись: ___________", size=9)
    _p(doc, "10. Разрешаю допуск к выполнению работ: _________ / __________________ /", size=9)

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    return out
