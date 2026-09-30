"""DOCX generator for the Belarus occupational-safety prescription form."""
from datetime import date, timedelta
from io import BytesIO
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from config import (
    PRESCRIPTION_DEFAULT_DEADLINE_DAYS,
    PRESCRIPTION_ORGANIZATION,
    PRESCRIPTION_SERVICE_NAME,
    PRESCRIPTION_DEFAULT_RESPONSIBLE,
)


def _font(run, size=9, bold=False, italic=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _cell(cell, text="", size=8, bold=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text or ""))
    _font(r, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _width(cell, cm):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(cm * 567)))
    tcW.set(qn("w:type"), "dxa")


def _borders(table, color="000000", size="4"):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = borders.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge)
            borders.append(e)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), size)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tblPr.append(borders)


def _basis(item):
    result = []
    for x in item.get("legal_basis") or []:
        doc = str(x.get("document") or "").strip()
        point = str(x.get("point") or "").strip()
        if doc:
            result.append(doc + (f", {point}" if point else ""))
    return "; ".join(result)


def build_draft_prescription(
    findings: List[Dict[str, Any]],
    photo_bytes: Optional[bytes] = None,
    enterprise: Optional[str] = None,
    subdivision: str = "",
    workplace: str = "",
    responsible: Optional[str] = None,
    recipient: str = "",
    prescription_number: str = "",
    issue_date: Optional[str] = None,
    deadline: Optional[str] = None,
    issuer_name: str = "",
    issuer_position: str = "",
    recipient_name: str = "",
    recipient_position: str = "",
) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.2)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(9)

    enterprise = enterprise or PRESCRIPTION_ORGANIZATION
    responsible = responsible or PRESCRIPTION_DEFAULT_RESPONSIBLE
    today = issue_date or date.today().strftime("%d.%m.%Y")
    due = deadline or (
        date.today() + timedelta(days=PRESCRIPTION_DEFAULT_DEADLINE_DAYS)
    ).strftime("%d.%m.%Y")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(PRESCRIPTION_SERVICE_NAME)
    _font(r, 8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(enterprise)
    _font(r, 8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"ПРЕДПИСАНИЕ № {prescription_number or '____________'}")
    _font(r, 13, True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"От {today} г.")
    _font(r, 9)

    for label, value in [
        ("Кому ", recipient or "____________________________________________"),
        ("Подразделение ", subdivision or "____________________________________________"),
        ("Рабочее место (объект) ", workplace or "____________________________________________"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label)
        _font(r, 9, True)
        r = p.add_run(value)
        _font(r, 9)

    p = doc.add_paragraph()
    r = p.add_run("(фамилия, имя, отчество, должность уполномоченного должностного лица работодателя, которому вручается предписание)")
    _font(r, 7, italic=True)

    p = doc.add_paragraph()
    r = p.add_run(
        "В соответствии со статьей 17 Закона Республики Беларусь от 23 июня 2008 года "
        "«Об охране труда» предлагаю выполнить следующие мероприятия:"
    )
    _font(r, 9)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)
    widths = [0.8, 11.0, 3.0, 3.0]
    headers = [
        "№\nп/п",
        "Перечень мероприятий (излагаются в виде требования со ссылкой на статьи, пункты "
        "(их части, подпункты и тому подобное) соответствующих нормативных правовых актов, "
        "технических нормативных правовых актов)",
        "Срок выполнения\n(дата)",
        "Отметка о выполнении\n(дата)",
    ]
    for c, h, w in zip(table.rows[0].cells, headers, widths):
        _width(c, w)
        _cell(c, h, 7, True, WD_ALIGN_PARAGRAPH.CENTER)

    confirmed = [
        x for x in findings
        if x.get("status") == "confirmed" and x.get("legal_basis")
    ]
    if confirmed:
        for i, item in enumerate(confirmed, 1):
            cells = table.add_row().cells
            for c, w in zip(cells, widths):
                _width(c, w)
            action = str(item.get("corrective_action") or "Устранить выявленное нарушение требований охраны труда.").strip()
            violation = str(item.get("violation") or "").strip()
            basis = _basis(item)
            measure = action
            if basis:
                measure += f" Ссылка на НПА/ТНПА: {basis}."
            if violation:
                measure += f"\nВыявлено: {violation}."
            _cell(cells[0], i, 8, False, WD_ALIGN_PARAGRAPH.CENTER)
            _cell(cells[1], measure, 8)
            _cell(cells[2], due, 8, False, WD_ALIGN_PARAGRAPH.CENTER)
            _cell(cells[3], "", 8)
    else:
        cells = table.add_row().cells
        for c, w in zip(cells, widths):
            _width(c, w)
        _cell(cells[0], "—", 8, False, WD_ALIGN_PARAGRAPH.CENTER)
        _cell(cells[1], "Подтвержденные нарушения с конкретным нормативным основанием отсутствуют. Документ требует проверки специалистом.", 8)
        _cell(cells[2], "", 8)
        _cell(cells[3], "", 8)

    for text in [
        "О выполнении мероприятий настоящего предписания по истечении указанных в нем сроков прошу письменно сообщить в службу охраны труда (специалисту по охране труда).",
        "Эксплуатация оборудования, инструмента, приспособлений, транспортных средств, выполнение работ (оказание услуг), которые были приостановлены (запрещены) в связи с угрозой для жизни или здоровья работающих и окружающих, согласно пунктам __________ настоящего предписания могут быть возобновлены после устранения выявленных нарушений требований по охране труда с разрешения службы охраны труда (специалиста по охране труда).",
        "Руководитель организации, уполномоченное должностное лицо работодателя, самовольно допустившие эксплуатацию оборудования, инструмента, приспособлений, транспортных средств, выполнение работ (оказание услуг), которые были приостановлены (запрещены) предписанием, привлекаются к ответственности в соответствии с законодательством.",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        r = p.add_run(text)
        _font(r, 8)

    sig = doc.add_table(rows=2, cols=2)
    sig.autofit = False
    _no_borders(sig)
    for row in sig.rows:
        _width(row.cells[0], 8.5)
        _width(row.cells[1], 9.0)
    _cell(sig.cell(0, 0), "Предписание выдал\n\n__________________________\n(подпись)", 8)
    _cell(sig.cell(0, 1), "Предписание получил\n\n__________________________\n(подпись)", 8)
    _cell(sig.cell(1, 0), f"{issuer_name or 'И.О. Фамилия'}\n{issuer_position or 'должность работника службы охраны труда (специалиста по охране труда)'}\nДата: {today}", 7)
    _cell(sig.cell(1, 1), f"{recipient_name or 'И.О. Фамилия'}\n{recipient_position or 'должность'}\nДата: __________________", 7)

    p = doc.add_paragraph()
    r = p.add_run("Примечания:")
    _font(r, 8, True)
    for n in [
        "1. Предписание составляется в 2 экземплярах, один из которых по принадлежности выдается уполномоченному должностному лицу работодателя, второй – остается в службе охраны труда (у специалиста по охране труда).",
        "2. Предписание регистрируется в службе охраны труда (у специалиста по охране труда) и хранится в течение 5 лет.",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.3)
        r = p.add_run(n)
        _font(r, 7)

    if photo_bytes:
        doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("ПРИЛОЖЕНИЕ — ФОТОГРАФИЧЕСКОЕ ПОДТВЕРЖДЕНИЕ")
        _font(r, 10, True)
        try:
            doc.add_picture(BytesIO(photo_bytes), width=Cm(15.5))
        except Exception:
            pass

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
