import argparse
import logging
import os
import re
import sys
import time
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple
from uuid import uuid4

from docx import Document
from dotenv import load_dotenv
from supabase import create_client

from embedding import (
    EMBEDDING_DIM,
    EMBEDDING_MODEL,
    get_document_embeddings,
    warmup_model,
)


# ============================================================
# CONFIG
# ============================================================

load_dotenv()

SCRIPT_DIR = Path(__file__).resolve().parent

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_SERVICE_ROLE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")

REINDEX_ALL = (
    os.getenv("REINDEX_ALL", "false").lower() == "true"
)

REINDEX_DOCS = {
    item.strip()
    for item in os.getenv("REINDEX_DOCS", "").split(",")
    if item.strip()
}

EMBEDDING_BATCH_SIZE = max(
    1,
    int(os.getenv("EMBEDDING_BATCH_SIZE", "8")),
)

UPLOAD_BATCH_SIZE = max(
    1,
    int(os.getenv("UPLOAD_BATCH_SIZE", "50")),
)

SUPABASE_RETRIES = max(
    1,
    int(os.getenv("SUPABASE_RETRIES", "3")),
)

SUPABASE_RETRY_DELAY = max(
    0.0,
    float(os.getenv("SUPABASE_RETRY_DELAY", "2")),
)

# Временное имя для staging-загрузки.
# Оно никогда не используется поиском ассистента.
STAGING_PREFIX = "__NPA_STAGING__"


# ============================================================
# LOGGING
# ============================================================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
)

logger = logging.getLogger("npa_loader")


# ============================================================
# DOCUMENT MAPPING
# ============================================================

DOC_NAME_MAP: Dict[str, str] = {

    "Закон об охране труда от 23 июня 2008 г. № 356-З":
        "Закон об охране труда от 23 июня 2008 г. № 356-З",

    "Кодекс Республики Беларусь об административных правонарушениях от 6 января 2021 г. № 91-З":
        "Кодекс Республики Беларусь об административных правонарушениях от 6 января 2021 г. № 91-З",

    "О бесплатном обеспечении работников молоком или равноценными пищевыми продуктами при работе с вредными веществами 27 февраля 2002 г. № 260":
        "О бесплатном обеспечении работников молоком или равноценными пищевыми продуктами при работе с вредными веществами 27 февраля 2002 г. № 260",

    "О документах, необходимых для расследования и учета несчастных случаев на производстве и профессиональных заболеваний Минтруда и Соцзащиты от 4 октября 2024 г. № 81 144":
        "О документах, необходимых для расследования и учета несчастных случаев на производстве и профессиональных заболеваний Минтруда и Соцзащиты от 4 октября 2024 г. № 81 144",

    "О контроле состояния водителей от 9 июля 2013 г. № 25.28":
        "О контроле состояния водителей от 9 июля 2013 г. № 25.28",

    "О мерах по укреплению общественной безопасности и дисциплины Директива от 11 марта 2004 г. № 1":
        "О мерах по укреплению общественной безопасности и дисциплины Директива от 11 марта 2004 г. № 1",

    "О пожарной безопасности Закон РБ от 15 июня 1993 г. № 2403-XII":
        "О пожарной безопасности Закон РБ от 15 июня 1993 г. № 2403-XII",

    "О порядке обучения, стажировки, инструктажа и проверки знаний работающих по вопросам охраны труда № 175 от 28 ноября 2008 г":
        "О порядке обучения, стажировки, инструктажа и проверки знаний работающих по вопросам охраны труда № 175 от 28 ноября 2008 г",

    "О порядке проведения предрейсовых и иных медицинских обследований водителей механических транспортных средств (за исключением колесных тракторов) от 3 декабря 2002 г. № 84":
        "О порядке проведения предрейсовых и иных медицинских обследований водителей механических транспортных средств (за исключением колесных тракторов) от 3 декабря 2002 г. № 84",

    "О порядке разработки и принятия локальных правовых актов по охране труда от 28 ноября 2008 г. № 176":
        "О порядке разработки и принятия локальных правовых актов по охране труда от 28 ноября 2008 г. № 176",

    "О порядке расследования и учета несчастных случаев МЧС от 6 января 2023 г. № 6":
        "О порядке расследования и учета несчастных случаев МЧС от 6 января 2023 г. № 6",

    "О проведении обязательных и внеочередных медицинских осмотров работающих от 29 июля 2019 г. № 74":
        "О проведении обязательных и внеочередных медицинских осмотров работающих от 29 июля 2019 г. № 74",

    "О расследовании и учете несчастных случаев на производстве и профессиональных заболеваний от 15 января 2004 г. № 30":
        "О расследовании и учете несчастных случаев на производстве и профессиональных заболеваний от 15 января 2004 г. № 30",

    "Об обеспечении пожарной безопасности постановление МЧС 21 декабря 2021 г. № 82":
        "Об обеспечении пожарной безопасности постановление МЧС 21 декабря 2021 г. № 82",

    "Об утверждении Инструкции о порядке осуществления контроля за соблюдением работниками требований по охране труда от 15 мая 2020 г. № 51":
        "Об утверждении Инструкции о порядке осуществления контроля за соблюдением работниками требований по охране труда от 15 мая 2020 г. № 51",

    "Об утверждении Межотраслевых правил по охране труда при проведении погрузочно-разгрузочных работ от 26 января 2018 г. № 12":
        "Об утверждении Межотраслевых правил по охране труда при проведении погрузочно-разгрузочных работ от 26 января 2018 г. № 12",

    "Об утверждении Межотраслевых правил по охране труда при эксплуатации напольного безрельсового транспорта и грузовых тележек от 30 декабря 2003 г. № 165":
        "Об утверждении Межотраслевых правил по охране труда при эксплуатации напольного безрельсового транспорта и грузовых тележек от 30 декабря 2003 г. № 165",

    "Об утверждении Правил по охране труда при выполнении работ на высоте от 6 февраля 2025 г. № 11":
        "Об утверждении Правил по охране труда при выполнении работ на высоте от 6 февраля 2025 г. № 11",

    "Об утверждении Правил по охране труда при выполнении строительных работ от 31 мая 2019 г. № 24 33":
        "Об утверждении Правил по охране труда при выполнении строительных работ от 31 мая 2019 г. № 24 33",

    "Об утверждении Правил по охране труда при производстве пищевой продукции от 31 декабря 2024 г. № 122":
        "Об утверждении Правил по охране труда при производстве пищевой продукции от 31 декабря 2024 г. № 122",

    "Об утверждении Правил по охране труда при эксплуатации автомобильного и городского электрического транспорта от 6 декабря 2022 г. № 78 104":
        "Об утверждении Правил по охране труда при эксплуатации автомобильного и городского электрического транспорта от 6 декабря 2022 г. № 78 104",

    "Об утверждении специфических санитарно-эпидемиологических требований от 24 января 2020 г. № 42":
        "Об утверждении специфических санитарно-эпидемиологических требований от 24 января 2020 г. № 42",

    "Об утверждении специфических санитарно-эпидемиологических требований Постановление от 1 февраля 2020 г. № 66":
        "Об утверждении специфических санитарно-эпидемиологических требований Постановление от 1 февраля 2020 г. № 66",

    "Об утверждении специфических требований по обеспечению пожарной безопасности взрывопожароопасных и пожароопасных производств 20 ноября 2019 г. № 779":
        "Об утверждении специфических требований по обеспечению пожарной безопасности взрывопожароопасных и пожароопасных производств 20 ноября 2019 г. № 779",

    "Правила по обеспечению СИЗ №209":
        "Правила по обеспечению СИЗ №209",

    "Правила по охране труда № 53":
        "Правила по охране труда № 53",

    "Трудовой кодекс Республики Беларусь 2026":
        "Трудовой кодекс Республики Беларусь 2026",

    "Об утверждении Инструкции по оценке условий труда при аттестации рабочих мест по условиям труда от 22 февраля 2008 г. № 35":
        "Об утверждении Инструкции по оценке условий труда при аттестации рабочих мест по условиям труда от 22 февраля 2008 г. № 35",

    "Об аттестации рабочих мест по условиям труда от 22 февраля 2008 г. № 253":
        "Об аттестации рабочих мест по условиям труда от 22 февраля 2008 г. № 253",

    "О подготовке и проверке знаний по вопросам промышленной безопасности от 6 июля 2016 г. № 31":
        "О подготовке и проверке знаний по вопросам промышленной безопасности от 6 июля 2016 г. № 31",

    "Закон РБ О промышленной безопасности от 5 января 2016 г. № 354-З":
        "Закон РБ О промышленной безопасности от 5 января 2016 г. № 354-З",
}


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text: str) -> str:
    text = str(text or "")
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_filename_stem(
    file_path: Path,
) -> str:
    return file_path.stem.strip()


# ============================================================
# DOCUMENT NAME RESOLUTION
# ============================================================

def resolve_doc_name(
    file_path: Path,
) -> Optional[str]:
    stem = normalize_filename_stem(file_path)

    if stem in DOC_NAME_MAP:
        return DOC_NAME_MAP[stem]

    normalized_stem = normalize_text(stem)

    for key, value in DOC_NAME_MAP.items():
        if normalize_text(key) == normalized_stem:
            return value

    return None


def resolve_explicit_doc_name(
    doc_name: str,
) -> Optional[str]:
    """
    Разрешает явное указание канонического имени документа.

    Это позволяет заменить редакцию, даже если новое имя DOCX
    отличается от старого имени файла.
    """

    requested = normalize_text(doc_name)

    if not requested:
        return None

    # Пользователь может передать имя DOCX с расширением.
    # В БД каноническое имя хранится без .docx.
    if requested.lower().endswith(".docx"):
        requested = requested[:-5].rstrip()

    if requested in DOC_NAME_MAP.values():
        return requested

    for key, value in DOC_NAME_MAP.items():
        if normalize_text(key) == requested:
            return value

    # Если имя явно передано пользователем,
    # считаем его каноническим.
    return requested


# ============================================================
# LEGAL DOMAIN / TOPIC
# ============================================================

def get_document_classification(
    doc_name: str,
) -> Tuple[str, str]:

    name = normalize_text(doc_name).lower()

    # --------------------------------------------------------
    # АТТЕСТАЦИЯ РАБОЧИХ МЕСТ
    # --------------------------------------------------------

    if (
        "аттестации рабочих мест" in name
        or "оценке условий труда при аттестации" in name
    ):
        return (
            "occupational_safety",
            "workplace_attestation",
        )

    # --------------------------------------------------------
    # ПОЖАРНАЯ БЕЗОПАСНОСТЬ
    # --------------------------------------------------------

    if "пожар" in name:
        return (
            "fire_safety",
            "fire_safety",
        )

    # --------------------------------------------------------
    # САНИТАРНЫЕ ТРЕБОВАНИЯ
    # --------------------------------------------------------

    if "санитар" in name:
        return (
            "sanitary",
            "sanitary_requirements",
        )

    # --------------------------------------------------------
    # ОХРАНА ТРУДА
    # --------------------------------------------------------

    if any(
        keyword in name
        for keyword in [
            "охране труда",
            "охрана труда",
            "сиз",
            "средств индивидуальной защиты",
            "несчастных случаев",
            "медицинских осмотров",
            "работ на высоте",
            "погрузочно-разгрузочных",
        ]
    ):
        return (
            "occupational_safety",
            "occupational_safety",
        )

    # --------------------------------------------------------
    # ПРОМЫШЛЕННАЯ БЕЗОПАСНОСТЬ
    # --------------------------------------------------------

    if any(
        keyword in name
        for keyword in [
            "промышленной безопасности",
            "промышленная безопасность",
            "опасных производственных объектов",
            "опасные производственные объекты",
            "опасный производственный объект",
            "производственный контроль",
            "техническое расследование аварий",
            "экспертизы промышленной безопасности",
            "экспертиза промышленной безопасности",
        ]
    ):
        return (
            "industrial_safety",
            "industrial_safety",
        )

    # --------------------------------------------------------
    # ОБЩИЙ ДОМЕН
    # --------------------------------------------------------

    return (
        "general",
        "general",
    )


# ============================================================
# DOCX READER
# ============================================================

def read_docx(
    file_path: Path,
) -> str:

    logger.info(
        "READ DOCX | %s",
        file_path.name,
    )

    document = Document(
        str(file_path)
    )

    parts: List[str] = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            cells: List[str] = []

            for cell in row.cells:

                cell_parts = []

                for paragraph in cell.paragraphs:

                    paragraph_text = (
                        paragraph.text.strip()
                    )

                    if paragraph_text:
                        cell_parts.append(
                            paragraph_text
                        )

                cell_text = " ".join(
                    cell_parts
                )

                if cell_text:
                    cells.append(
                        cell_text
                    )

            if cells:
                parts.append(
                    " | ".join(cells)
                )

    text = "\n".join(parts)

    text = text.replace(
        "\r\n",
        "\n",
    )

    text = text.replace(
        "\r",
        "\n",
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


# ============================================================
# CHUNKING
# ============================================================

def split_text_into_chunks(
    text: str,
) -> List[Tuple[str, str]]:

    text = str(
        text or ""
    ).strip()

    if not text:
        return []

    paragraphs = [
        paragraph.strip()
        for paragraph in re.split(
            r"\n+",
            text,
        )
        if paragraph.strip()
    ]

    if not paragraphs:
        return []

    chunks: List[
        Tuple[str, str]
    ] = []

    current_point = "general"

    current_lines: List[str] = []

    marker_pattern = re.compile(
        r"^(Статья\s+\d+(?:\.\d+)*"
        r"|Пункт\s+\d+(?:\.\d+)*"
        r"|\d+(?:\.\d+)*\.)",
        re.IGNORECASE,
    )

    def flush_current() -> None:

        nonlocal current_lines

        if not current_lines:
            return

        content = "\n".join(
            current_lines
        ).strip()

        if content:

            chunks.append(
                (
                    current_point,
                    content,
                )
            )

        current_lines = []

    for paragraph in paragraphs:

        match = marker_pattern.match(
            paragraph
        )

        if match:

            flush_current()

            current_point = (
                match.group(1).strip()
            )

            current_lines = [
                paragraph
            ]

        else:

            current_lines.append(
                paragraph
            )

    flush_current()

    return chunks


# ============================================================
# SOURCE FILE DISCOVERY
# ============================================================

def find_source_files() -> List[Path]:

    files: List[Path] = []

    for path in SCRIPT_DIR.iterdir():

        if not path.is_file():
            continue

        if path.name.startswith("~$"):
            continue

        if path.suffix.lower() != ".docx":
            continue

        files.append(path)

    return sorted(
        files,
        key=lambda item: item.name.lower(),
    )


# ============================================================
# SOURCE VALIDATION
# ============================================================

def validate_source_files(
    files: Sequence[Path],
    explicit_doc_name: Optional[str] = None,
) -> None:

    errors: List[str] = []

    for file_path in files:

        doc_name = (
            explicit_doc_name
            if explicit_doc_name
            else resolve_doc_name(file_path)
        )

        if not doc_name:

            errors.append(
                f"{file_path.name}: "
                f"нет записи в DOC_NAME_MAP. "
                f"Нормализованное имя: "
                f"'{normalize_filename_stem(file_path)}'"
            )

    if errors:

        logger.error("")

        logger.error(
            "ОШИБКИ СОПОСТАВЛЕНИЯ ФАЙЛОВ"
        )

        for error in errors:
            logger.error(
                " - %s",
                error,
            )

        raise RuntimeError(
            "Не все выбранные DOCX-файлы "
            "сопоставлены с DOC_NAME_MAP."
        )

    logger.info(
        "DOC_NAME_MAP | выбранные файлы "
        "успешно сопоставлены."
    )


# ============================================================
# SUPABASE CLIENT
# ============================================================

def get_supabase_client():

    if not SUPABASE_URL:

        raise RuntimeError(
            "SUPABASE_URL не задан."
        )

    if not SUPABASE_SERVICE_ROLE_KEY:

        raise RuntimeError(
            "SUPABASE_SERVICE_ROLE_KEY не задан."
        )

    return create_client(
        SUPABASE_URL,
        SUPABASE_SERVICE_ROLE_KEY,
    )


# ============================================================
# SUPABASE RETRY
# ============================================================

def supabase_execute(
    operation,
    description: str = "Supabase operation",
):

    last_error = None

    for attempt in range(
        1,
        SUPABASE_RETRIES + 1,
    ):

        try:

            return operation()

        except Exception as exc:

            last_error = exc

            logger.warning(
                "%s | attempt %s/%s failed: %s",
                description,
                attempt,
                SUPABASE_RETRIES,
                exc,
            )

            if (
                attempt
                < SUPABASE_RETRIES
            ):

                time.sleep(
                    SUPABASE_RETRY_DELAY
                    * attempt
                )

    raise RuntimeError(
        f"{description} failed after "
        f"{SUPABASE_RETRIES} attempts: "
        f"{last_error}"
    )


# ============================================================
# SUPABASE CONNECTION
# ============================================================

def test_supabase_connection(
    supabase,
) -> None:

    def operation():

        return (
            supabase
            .table("npa_chunks")
            .select("id")
            .limit(1)
            .execute()
        )

    supabase_execute(
        operation,
        "SUPABASE connection test",
    )

    logger.info(
        "SUPABASE | соединение успешно."
    )


# ============================================================
# DOCUMENT EXISTENCE
# ============================================================

def document_exists(
    supabase,
    doc_name: str,
) -> bool:

    def operation():

        return (
            supabase
            .table("npa_chunks")
            .select("id")
            .eq(
                "doc_name",
                doc_name,
            )
            .limit(1)
            .execute()
        )

    response = supabase_execute(
        operation,
        f"DOCUMENT EXISTS | {doc_name}",
    )

    return bool(response.data)


# ============================================================
# DOCUMENT ID
# ============================================================

def get_document_id(
    supabase,
    doc_name: str,
) -> str:
    """
    Возвращает UUID записи из npa_documents.

    npa_chunks.document_id является обязательным FK,
    поэтому каждый chunk существующего НПА должен
    ссылаться на соответствующую запись npa_documents.
    """

    def operation():

        return (
            supabase
            .table("npa_documents")
            .select("id")
            .eq(
                "doc_name",
                doc_name,
            )
            .limit(1)
            .execute()
        )

    response = supabase_execute(
        operation,
        f"DOCUMENT ID | {doc_name}",
    )

    data = getattr(
        response,
        "data",
        None,
    ) or []

    if not data:
        raise RuntimeError(
            "В npa_documents не найдена запись "
            f"для документа: {doc_name}"
        )

    document_id = data[0].get("id")

    if not document_id:
        raise RuntimeError(
            "В npa_documents отсутствует id "
            f"для документа: {doc_name}"
        )

    logger.info(
        "DOCUMENT ID | %s | %s",
        doc_name,
        document_id,
    )

    return str(document_id)


# ============================================================
# DOCUMENT COUNT
# ============================================================

def count_document_rows(
    supabase,
    doc_name: str,
) -> int:
    """
    Возвращает фактическое количество chunks документа.
    """

    def operation():

        return (
            supabase
            .table("npa_chunks")
            .select(
                "id",
                count="exact",
            )
            .eq(
                "doc_name",
                doc_name,
            )
            .limit(1)
            .execute()
        )

    response = supabase_execute(
        operation,
        f"COUNT DOCUMENT | {doc_name}",
    )

    count = getattr(
        response,
        "count",
        None,
    )

    if count is None:

        data = getattr(
            response,
            "data",
            None,
        )

        return len(
            data or []
        )

    return int(count)


# ============================================================
# REINDEX CONFIG
# ============================================================

def should_reindex_document(
    doc_name: str,
) -> bool:

    if REINDEX_ALL:
        return True

    if (
        REINDEX_DOCS
        and doc_name in REINDEX_DOCS
    ):
        return True

    return False


# ============================================================
# DELETE DOCUMENT
# ============================================================

def delete_document_vectors(
    supabase,
    doc_name: str,
) -> None:

    logger.info(
        "DELETE | chunks документа: %s",
        doc_name,
    )

    def operation():

        return (
            supabase
            .table("npa_chunks")
            .delete()
            .eq(
                "doc_name",
                doc_name,
            )
            .execute()
        )

    supabase_execute(
        operation,
        f"DELETE DOCUMENT | {doc_name}",
    )


# ============================================================
# RENAME STAGING DOCUMENT
# ============================================================

def update_document_name(
    supabase,
    old_doc_name: str,
    new_doc_name: str,
) -> None:
    """
    Переводит staging chunks в каноническое имя.
    """

    logger.info(
        "ACTIVATE | %s -> %s",
        old_doc_name,
        new_doc_name,
    )

    def operation():

        return (
            supabase
            .table("npa_chunks")
            .update(
                {
                    "doc_name": new_doc_name
                }
            )
            .eq(
                "doc_name",
                old_doc_name,
            )
            .execute()
        )

    supabase_execute(
        operation,
        (
            "ACTIVATE DOCUMENT | "
            f"{old_doc_name} -> {new_doc_name}"
        ),
    )


# ============================================================
# EMBEDDING VALIDATION
# ============================================================

def validate_embeddings(
    embeddings: Sequence[
        Sequence[float]
    ],
    expected_count: int,
) -> None:

    if (
        len(embeddings)
        != expected_count
    ):

        raise RuntimeError(
            "Embedding count mismatch: "
            f"expected {expected_count}, "
            f"got {len(embeddings)}"
        )

    for index, embedding in enumerate(
        embeddings
    ):

        if (
            len(embedding)
            != EMBEDDING_DIM
        ):

            raise RuntimeError(
                "Embedding dimension mismatch "
                f"at index {index}: "
                f"expected {EMBEDDING_DIM}, "
                f"got {len(embedding)}"
            )


def validate_embedding_model() -> None:

    logger.info(
        "EMBEDDING MODEL | %s",
        EMBEDDING_MODEL,
    )

    logger.info(
        "EMBEDDING DIM | %s",
        EMBEDDING_DIM,
    )

    warmup_model()

    test_embeddings = (
        get_document_embeddings(
            [
                "Тестовый фрагмент "
                "нормативного правового "
                "акта по охране труда."
            ]
        )
    )

    validate_embeddings(
        test_embeddings,
        1,
    )

    logger.info(
        "EMBEDDING TEST | OK | dimension=%s",
        len(
            test_embeddings[0]
        ),
    )


# ============================================================
# PREPARE CHUNKS
# ============================================================

def prepare_chunks(
    doc_name: str,
    chunks: Sequence[
        Tuple[str, str]
    ],
    legal_domain: str,
    topic: str,
    storage_doc_name: Optional[str] = None,
) -> List[dict]:
    """
    Формирует уникальные chunks.

    storage_doc_name используется для staging.
    """

    actual_doc_name = (
        storage_doc_name
        or doc_name
    )

    prepared_chunks: List[dict] = []

    seen_chunks = set()

    duplicate_chunks = 0

    for point_num, content in chunks:

        content = content.strip()

        if not content:
            continue

        chunk_key = (
            str(point_num).strip(),
            content,
        )

        if chunk_key in seen_chunks:

            duplicate_chunks += 1

            continue

        seen_chunks.add(
            chunk_key
        )

        prepared_chunks.append(
            {
                "doc_name": actual_doc_name,
                "doc_type": "НПА",
                "point_num": point_num,
                "content": content,
                "legal_domain": legal_domain,
                "topic": topic,
            }
        )

    if duplicate_chunks:

        logger.warning(
            "CHUNKS | удалено внутренних дублей=%s",
            duplicate_chunks,
        )

    logger.info(
        "CHUNKS | unique to upload=%s",
        len(prepared_chunks),
    )

    return prepared_chunks


# ============================================================
# GENERATE EMBEDDINGS
# ============================================================

def generate_embeddings_for_chunks(
    prepared_chunks: Sequence[dict],
) -> List[List[float]]:

    if not prepared_chunks:
        return []

    all_embeddings: List[
        List[float]
    ] = []

    total_embedding_batches = (
        len(prepared_chunks)
        + EMBEDDING_BATCH_SIZE
        - 1
    ) // EMBEDDING_BATCH_SIZE

    for start in range(
        0,
        len(prepared_chunks),
        EMBEDDING_BATCH_SIZE,
    ):

        batch = prepared_chunks[
            start:
            start + EMBEDDING_BATCH_SIZE
        ]

        batch_number = (
            start
            // EMBEDDING_BATCH_SIZE
        ) + 1

        logger.info(
            "EMBEDDING | batch %s/%s | chunks=%s",
            batch_number,
            total_embedding_batches,
            len(batch),
        )

        batch_texts = [
            item["content"]
            for item in batch
        ]

        embeddings = (
            get_document_embeddings(
                batch_texts
            )
        )

        validate_embeddings(
            embeddings,
            len(batch),
        )

        all_embeddings.extend(
            embeddings
        )

    validate_embeddings(
        all_embeddings,
        len(prepared_chunks),
    )

    return all_embeddings


# ============================================================
# BUILD SUPABASE ROWS
# ============================================================

def build_rows(
    prepared_chunks: Sequence[dict],
    embeddings: Sequence[
        Sequence[float]
    ],
) -> List[dict]:

    validate_embeddings(
        embeddings,
        len(prepared_chunks),
    )

    rows: List[dict] = []

    for item, embedding in zip(
        prepared_chunks,
        embeddings,
    ):

        rows.append(
            {
                "doc_name": item["doc_name"],
                "doc_type": item["doc_type"],
                "point_num": item["point_num"],
                "content": item["content"],
                "legal_domain": item["legal_domain"],
                "topic": item["topic"],
                "embedding": embedding,
            }
        )

    return rows


# ============================================================
# SUPABASE INSERT
# ============================================================

def insert_rows(
    supabase,
    rows: List[dict],
) -> int:

    if not rows:
        return 0

    uploaded = 0

    total_batches = (
        len(rows)
        + UPLOAD_BATCH_SIZE
        - 1
    ) // UPLOAD_BATCH_SIZE

    for start in range(
        0,
        len(rows),
        UPLOAD_BATCH_SIZE,
    ):

        batch = rows[
            start:
            start + UPLOAD_BATCH_SIZE
        ]

        batch_number = (
            start
            // UPLOAD_BATCH_SIZE
        ) + 1

        logger.info(
            "SUPABASE INSERT | batch %s/%s | rows=%s",
            batch_number,
            total_batches,
            len(batch),
        )

        def operation():

            return (
                supabase
                .table("npa_chunks")
                .insert(batch)
                .execute()
            )

        supabase_execute(
            operation,
            (
                f"INSERT BATCH "
                f"{batch_number}/{total_batches}"
            ),
        )

        uploaded += len(batch)

    return uploaded


# ============================================================
# STAGING CLEANUP
# ============================================================

def cleanup_staging(
    supabase,
    staging_doc_name: str,
) -> None:

    try:

        count = count_document_rows(
            supabase,
            staging_doc_name,
        )

        if count <= 0:
            return

        logger.warning(
            "STAGING CLEANUP | "
            "удаляем %s временных chunks: %s",
            count,
            staging_doc_name,
        )

        delete_document_vectors(
            supabase,
            staging_doc_name,
        )

    except Exception as exc:

        logger.error(
            "STAGING CLEANUP FAILED | %s | %s",
            staging_doc_name,
            exc,
        )


# ============================================================
# SAFE REPLACEMENT
# ============================================================

def replace_document_safely(
    supabase,
    doc_name: str,
    rows: List[dict],
    dry_run: bool = False,
) -> int:
    """
    Безопасная замена редакции НПА.

    Последовательность:

    1. Новая редакция полностью подготовлена.
    2. Embeddings рассчитаны.
    3. Rows сформированы.
    4. Rows загружаются под временным staging-именем.
    5. Staging проверяется.
    6. Только после этого удаляется старая редакция.
    7. Staging переименовывается в canonical doc_name.
    8. Выполняется финальная проверка.
    """

    if not rows:

        raise RuntimeError(
            f"Нет данных для замены документа: "
            f"{doc_name}"
        )

    expected_count = len(rows)

    if dry_run:

        logger.info(
            "DRY-RUN | замена НЕ выполняется."
        )

        logger.info(
            "DRY-RUN | документ=%s | "
            "новых chunks=%s",
            doc_name,
            expected_count,
        )

        return expected_count

    # --------------------------------------------------------
    # Проверяем, что старый документ действительно существует.
    # --------------------------------------------------------

    if not document_exists(
        supabase,
        doc_name,
    ):

        raise RuntimeError(
            "REPLACE требует существующий документ. "
            f"Документ '{doc_name}' "
            "не найден в Supabase."
        )

    # npa_chunks.document_id — обязательный FK.
    # Для replace используем тот же document_id, который
    # уже принадлежит существующей записи npa_documents.
    document_id = get_document_id(
        supabase,
        doc_name,
    )

    staging_doc_name = (
        f"{STAGING_PREFIX}"
        f"{uuid4().hex}"
    )

    staging_rows = []

    for row in rows:

        staged_row = dict(
            row
        )

        staged_row[
            "doc_name"
        ] = staging_doc_name

        staged_row[
            "document_id"
        ] = document_id

        staging_rows.append(
            staged_row
        )

    logger.info(
        "REPLACE | staging name=%s | document_id=%s",
        staging_doc_name,
        document_id,
    )

    try:

        # ====================================================
        # 1. ЗАГРУЖАЕМ НОВУЮ РЕДАКЦИЮ
        # ====================================================

        inserted = insert_rows(
            supabase,
            staging_rows,
        )

        if inserted != expected_count:

            raise RuntimeError(
                "STAGING INSERT COUNT MISMATCH: "
                f"expected {expected_count}, "
                f"inserted {inserted}"
            )

        # ====================================================
        # 2. ПРОВЕРЯЕМ STAGING
        # ====================================================

        staged_count = (
            count_document_rows(
                supabase,
                staging_doc_name,
            )
        )

        if staged_count != expected_count:

            raise RuntimeError(
                "STAGING VERIFY COUNT MISMATCH: "
                f"expected {expected_count}, "
                f"found {staged_count}"
            )

        logger.info(
            "STAGING VERIFY | OK | chunks=%s",
            staged_count,
        )

        # ====================================================
        # 3. ТЕПЕРЬ МОЖНО УДАЛИТЬ СТАРУЮ РЕДАКЦИЮ
        # ====================================================

        delete_document_vectors(
            supabase,
            doc_name,
        )

        old_count_after_delete = (
            count_document_rows(
                supabase,
                doc_name,
            )
        )

        if old_count_after_delete != 0:

            raise RuntimeError(
                "OLD DOCUMENT DELETE VERIFY FAILED: "
                f"remaining="
                f"{old_count_after_delete}"
            )

        # ====================================================
        # 4. АКТИВИРУЕМ НОВУЮ РЕДАКЦИЮ
        # ====================================================

        update_document_name(
            supabase,
            staging_doc_name,
            doc_name,
        )

        # ====================================================
        # 5. ФИНАЛЬНАЯ ПРОВЕРКА
        # ====================================================

        final_count = (
            count_document_rows(
                supabase,
                doc_name,
            )
        )

        if final_count != expected_count:

            raise RuntimeError(
                "FINAL DOCUMENT COUNT MISMATCH: "
                f"expected {expected_count}, "
                f"found {final_count}"
            )

        staging_left = (
            count_document_rows(
                supabase,
                staging_doc_name,
            )
        )

        if staging_left != 0:

            raise RuntimeError(
                "STAGING ROWS REMAIN "
                "AFTER ACTIVATION: "
                f"{staging_left}"
            )

        logger.info(
            "REPLACE COMPLETE | %s | chunks=%s",
            doc_name,
            final_count,
        )

        return final_count

    except Exception:

        # Удаляем только staging.
        # Канонический документ здесь не удаляем.
        cleanup_staging(
            supabase,
            staging_doc_name,
        )

        raise


# ============================================================
# ADD NEW DOCUMENT
# ============================================================

def add_document(
    supabase,
    doc_name: str,
    rows: List[dict],
    dry_run: bool = False,
) -> int:

    if not rows:

        raise RuntimeError(
            f"Нет данных для добавления документа: "
            f"{doc_name}"
        )

    expected_count = len(rows)

    if document_exists(
        supabase,
        doc_name,
    ):

        raise RuntimeError(
            "ADD отменён: документ уже "
            "существует в Supabase: "
            f"{doc_name}"
        )

    if dry_run:

        logger.info(
            "DRY-RUN | добавление НЕ выполняется."
        )

        logger.info(
            "DRY-RUN | документ=%s | chunks=%s",
            doc_name,
            expected_count,
        )

        return expected_count

    inserted = insert_rows(
        supabase,
        rows,
    )

    if inserted != expected_count:

        raise RuntimeError(
            "ADD INSERT COUNT MISMATCH: "
            f"expected {expected_count}, "
            f"inserted {inserted}"
        )

    actual_count = (
        count_document_rows(
            supabase,
            doc_name,
        )
    )

    if actual_count != expected_count:

        raise RuntimeError(
            "ADD VERIFY COUNT MISMATCH: "
            f"expected {expected_count}, "
            f"found {actual_count}"
        )

    logger.info(
        "ADD COMPLETE | %s | chunks=%s",
        doc_name,
        actual_count,
    )

    return actual_count


# ============================================================
# PROCESS FILE
# ============================================================

def process_file(
    supabase,
    file_path: Path,
    mode: str = "auto",
    explicit_doc_name: Optional[str] = None,
    dry_run: bool = False,
) -> Tuple[int, int]:
    """
    mode:

    auto:
        новый документ -> add
        существующий -> skip
        REINDEX -> safe replace

    add:
        только добавление

    replace:
        только безопасная замена
    """

    # --------------------------------------------------------
    # DOC NAME
    # --------------------------------------------------------

    if explicit_doc_name:

        doc_name = (
            resolve_explicit_doc_name(
                explicit_doc_name
            )
        )

    else:

        doc_name = resolve_doc_name(
            file_path
        )

    if not doc_name:

        raise RuntimeError(
            f"Не найден DOC_NAME_MAP "
            f"для '{file_path.name}'. "
            "Используйте --doc-name "
            "для явного указания "
            "канонического имени."
        )

    logger.info("")
    logger.info("=" * 70)

    logger.info(
        "PROCESS | %s",
        file_path.name,
    )

    logger.info(
        "DOC NAME | %s",
        doc_name,
    )

    # --------------------------------------------------------
    # CLASSIFICATION
    # --------------------------------------------------------

    legal_domain, topic = (
        get_document_classification(
            doc_name
        )
    )

    logger.info(
        "LEGAL DOMAIN | %s",
        legal_domain,
    )

    logger.info(
        "TOPIC | %s",
        topic,
    )

    logger.info("=" * 70)

    # --------------------------------------------------------
    # READ
    # --------------------------------------------------------

    text = read_docx(
        file_path
    )

    if not text:

        logger.warning(
            "SKIP | пустой документ: %s",
            file_path.name,
        )

        return 0, 0

    logger.info(
        "TEXT | chars=%s",
        len(text),
    )

    # --------------------------------------------------------
    # CHUNKS
    # --------------------------------------------------------

    chunks = (
        split_text_into_chunks(
            text
        )
    )

    if not chunks:

        logger.warning(
            "SKIP | chunks не найдены: %s",
            file_path.name,
        )

        return 0, 0

    logger.info(
        "CHUNKS | найдено: %s",
        len(chunks),
    )

    # --------------------------------------------------------
    # PREPARE
    # --------------------------------------------------------

    prepared_chunks = (
        prepare_chunks(
            doc_name=doc_name,
            chunks=chunks,
            legal_domain=legal_domain,
            topic=topic,
        )
    )

    if not prepared_chunks:

        raise RuntimeError(
            "После очистки не осталось chunks: "
            f"{doc_name}"
        )

    # --------------------------------------------------------
    # EMBEDDINGS
    #
    # КРИТИЧЕСКИ ВАЖНО:
    # здесь старая редакция ещё существует.
    # --------------------------------------------------------

    embeddings = (
        generate_embeddings_for_chunks(
            prepared_chunks
        )
    )

    # --------------------------------------------------------
    # BUILD ROWS
    # --------------------------------------------------------

    rows = build_rows(
        prepared_chunks,
        embeddings,
    )

    if len(rows) != len(
        prepared_chunks
    ):

        raise RuntimeError(
            "ROWS COUNT MISMATCH: "
            f"prepared="
            f"{len(prepared_chunks)}, "
            f"rows={len(rows)}"
        )

    # --------------------------------------------------------
    # EXPLICIT REPLACE
    # --------------------------------------------------------

    if mode == "replace":

        uploaded = (
            replace_document_safely(
                supabase,
                doc_name,
                rows,
                dry_run=dry_run,
            )
        )

        return uploaded, 0

    # --------------------------------------------------------
    # EXPLICIT ADD
    # --------------------------------------------------------

    if mode == "add":

        uploaded = add_document(
            supabase,
            doc_name,
            rows,
            dry_run=dry_run,
        )

        return uploaded, 0

    # --------------------------------------------------------
    # AUTO MODE
    # --------------------------------------------------------

    force_reindex = (
        should_reindex_document(
            doc_name
        )
    )

    if force_reindex:

        logger.info(
            "REINDEX | безопасная замена документа: %s",
            doc_name,
        )

        uploaded = (
            replace_document_safely(
                supabase,
                doc_name,
                rows,
                dry_run=dry_run,
            )
        )

        return uploaded, 0

    # --------------------------------------------------------
    # EXISTING
    # --------------------------------------------------------

    if document_exists(
        supabase,
        doc_name,
    ):

        logger.info(
            "SKIP EXISTING | документ уже "
            "есть в Supabase: %s",
            doc_name,
        )

        return 0, 1

    # --------------------------------------------------------
    # NEW DOCUMENT
    # --------------------------------------------------------

    logger.info(
        "NEW DOCUMENT | документ отсутствует "
        "в Supabase: %s",
        doc_name,
    )

    uploaded = add_document(
        supabase,
        doc_name,
        rows,
        dry_run=dry_run,
    )

    return uploaded, 0


# ============================================================
# CLI
# ============================================================

def parse_args() -> argparse.Namespace:

    parser = argparse.ArgumentParser(
        description=(
            "Загрузчик НПА Республики Беларусь "
            "в Supabase pgvector."
        )
    )

    mode_group = (
        parser.add_mutually_exclusive_group()
    )

    mode_group.add_argument(
        "--replace",
        metavar="FILE",
        help=(
            "Безопасно заменить существующую "
            "редакцию указанного НПА."
        ),
    )

    mode_group.add_argument(
        "--add",
        metavar="FILE",
        help=(
            "Добавить новый НПА. "
            "Если документ уже существует, "
            "операция будет отменена."
        ),
    )

    parser.add_argument(
        "--doc-name",
        help=(
            "Каноническое имя НПА в Supabase. "
            "Используется, если имя нового DOCX "
            "отличается от имени старого файла."
        ),
    )

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help=(
            "Прочитать DOCX, сформировать chunks "
            "и embeddings, но не изменять Supabase."