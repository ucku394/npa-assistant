"""Unit tests for the Belarus OHS Document Generator (Act H-1 & Work Permit)."""

from io import BytesIO
from docx import Document

from services.document_generator import (
    ActN1Request,
    WorkPermitRequest,
    auto_suggest_act_n1_rules,
    auto_suggest_work_permit_measures,
    generate_act_n1_docx,
    generate_work_permit_docx,
)


def test_auto_suggest_act_n1_height_incident():
    desc = "Работник упал с приставной лестницы при выполнении работ на высоте."
    suggested = auto_suggest_act_n1_rules(desc)

    combined_rules = " ".join(suggested["violated_npa"])
    assert "№ 11" in combined_rules
    assert "высоте" in combined_rules
    assert len(suggested["causes"]) > 0
    assert len(suggested["corrective_measures"]) > 0


def test_auto_suggest_act_n1_lathe_incident():
    desc = "При обработке вала на токарном станке деталь вылетела из патрона."
    suggested = auto_suggest_act_n1_rules(desc)

    combined_rules = " ".join(suggested["violated_npa"])
    assert "станок" in combined_rules or "металл" in combined_rules or "356" in combined_rules
    assert len(suggested["causes"]) > 0


def test_auto_suggest_work_permit_height():
    suggested = auto_suggest_work_permit_measures("Монтаж металлоконструкций на высоте")
    combined_ppe = " ".join(suggested["required_ppe"])
    assert "привязь" in combined_ppe.lower()
    assert "строп" in combined_ppe.lower()
    assert any("анкер" in m.lower() or "страховоч" in m.lower() for m in suggested["safety_measures"])


def test_auto_suggest_work_permit_hot_works():
    suggested = auto_suggest_work_permit_measures("Огневые сварочные работы на трубопроводе")
    combined_measures = " ".join(suggested["safety_measures"])
    assert "огнетушител" in combined_measures.lower()
    assert any("сварщик" in p.lower() or "щиток" in p.lower() for p in suggested["required_ppe"])


def test_generate_act_n1_docx():
    req = ActN1Request(
        organization_name="ОАО 'Тестовый Завод'",
        subdivision="Цех металлообработки",
        accident_location="Токарный станок 16К20",
        victim_fio="Петров Петр Петрович",
        victim_profession="Токарь 5 разряда",
        incident_description="При закреплении заготовки в патроне токарного станка произошло защемление пальца.",
    )
    docx_io = generate_act_n1_docx(req)
    assert isinstance(docx_io, BytesIO)
    content = docx_io.getvalue()
    # Check valid DOCX zip signature (PK\x03\x04)
    assert content.startswith(b"PK\x03\x04")

    # Read back with python-docx
    doc = Document(docx_io)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "АКТ №" in text
    assert "форма Н-1" in text
    assert "Петров Петр Петрович" in text
    assert "ОАО 'Тестовый Завод'" in text


def test_generate_work_permit_docx():
    req = WorkPermitRequest(
        organization_name="ЗАО 'СтройМонтаж'",
        subdivision="Участок фасадных работ",
        work_type="Работы на высоте по герметизации швов",
        work_location="Фасад здания, отметка +15.000",
        permit_issuer="Главный инженер Сидоров С.С.",
        work_leader="Прораб Васильев В.В.",
        work_supervisor="Бригадир Николаев Н.Н.",
        brigade_members=[
            {"fio": "Николаев Н.Н.", "qualification": "Бригадир / 2 группа"},
            {"fio": "Козлов К.К.", "qualification": "Монтажник / 1 группа"},
        ],
    )
    docx_io = generate_work_permit_docx(req)
    assert isinstance(docx_io, BytesIO)
    content = docx_io.getvalue()
    assert content.startswith(b"PK\x03\x04")

    doc = Document(docx_io)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "НАРЯД-ДОПУСК" in text
    assert "ЗАО 'СтройМонтаж'" in text
    assert "Сидоров С.С." in text
