"""Regression tests for DOCX block order and batched chunk renumbering."""

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

import scripts.upload as upload


class _FakeQuery:
    def __init__(self, client):
        self.client = client

    def upsert(self, rows, on_conflict=None):
        self.client.batches.append([dict(row) for row in rows])
        self.client.conflict_targets.append(on_conflict)
        return self

    def execute(self):
        return SimpleNamespace(data=[])


class _FakeSupabase:
    def __init__(self):
        self.batches = []
        self.conflict_targets = []

    def table(self, table_name):
        if table_name != "npa_chunks":
            raise AssertionError(f"Unexpected table: {table_name}")
        return _FakeQuery(self)


class UploadRegressionTests(unittest.TestCase):
    def test_read_docx_preserves_paragraph_table_paragraph_order(self):
        document = Document()
        document.add_paragraph("Пункт 12. Перед началом работы")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Параметр"
        table.cell(0, 1).text = "Значение"
        document.add_paragraph("Пункт 13. После таблицы")

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "table-order.docx"
            document.save(path)
            text = upload.read_docx(path)

        paragraph_before = text.index("Пункт 12.")
        table_position = text.index("Параметр | Значение")
        paragraph_after = text.index("Пункт 13.")
        self.assertLess(paragraph_before, table_position)
        self.assertLess(table_position, paragraph_after)

    def test_renumbering_uses_batched_upserts_and_sequential_numbers(self):
        supabase = _FakeSupabase()
        staged_rows = [
            {
                "id": f"row-{number}",
                "doc_name": "__NPA_STAGING__example",
                "document_id": "document-uuid",
                "chunk_no": 700 + number,
                "content": f"chunk-{number}",
                "embedding": [0.0],
            }
            for number in range(1, 124)
        ]

        with patch.object(upload, "UPLOAD_BATCH_SIZE", 50):
            upload.renumber_staged_rows(
                supabase,
                staged_rows,
                "Test document",
            )

        self.assertEqual([50, 50, 23], [len(batch) for batch in supabase.batches])
        self.assertEqual(["id", "id", "id"], supabase.conflict_targets)
        renumbered = [row for batch in supabase.batches for row in batch]
        self.assertEqual(list(range(1, 124)), [row["chunk_no"] for row in renumbered])
        self.assertTrue(
            all(row["doc_name"] == "__NPA_STAGING__example" for row in renumbered)
        )
        self.assertTrue(
            all(row["id"] == f"row-{index}" for index, row in enumerate(renumbered, 1))
        )


if __name__ == "__main__":
    unittest.main()
